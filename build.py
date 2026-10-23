"""Сборка: data.json → карточка-страница, прайс для Яндекс Бизнеса, картинки историй, логотип и обложка, тексты.
Запуск: python3 build.py   (нужны openpyxl и playwright для картинок)"""
import base64
import html
import json
import pathlib
import re
import shutil

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = pathlib.Path(__file__).parent
OUT = ROOT / "dist" / "euphoria"
IMP = OUT / "yandex-business"
data = json.loads((ROOT / "data.json").read_text(encoding="utf-8"))


def nbsp(v):  # цены не разрываются переносом: «3 300 ₽»
    if isinstance(v, str):
        return re.sub(r"(\d) (?=\d{3})", "\\1\u00a0", v).replace(" ₽", "\u00a0₽")
    if isinstance(v, list):
        return [nbsp(x) for x in v]
    if isinstance(v, dict):
        return {k: nbsp(x) for k, x in v.items()}
    return v


data = nbsp(data)
S = data["salon"]

if OUT.exists():
    shutil.rmtree(OUT)
(IMP / "stories").mkdir(parents=True)
(IMP / "brand-images").mkdir(parents=True)
# Фото салона для раздела «Фото» в Яндекс Бизнесе (обработанные снимки 19.10.2026)
shutil.copytree(ROOT / "photos" / "salon", IMP / "photos-salon")

# Шрифты встраиваются в страницу и картинки (работает без интернета). Лицензия SIL OFL.
RANGES = {"latin": "U+0000-00FF,U+0131,U+0152-0153,U+02BB-02BC,U+02C6,U+02DA,U+02DC,U+0304,U+0308,U+0329,U+2000-206F,U+20AC,U+2122,U+2191,U+2193,U+2212,U+2215,U+FEFF,U+FFFD",
          "cyrillic": "U+0301,U+0400-045F,U+0490-0491,U+04B0-04B1,U+2116"}
FAMILIES = {"comfortaa": "Comfortaa", "montserrat": "Montserrat", "bodoni-moda": "Bodoni Moda"}
faces = []
for f in sorted((ROOT / "fonts").glob("*.woff2")):
    stem = f.stem  # comfortaa-cyrillic-300-normal
    for key, fam in FAMILIES.items():
        if stem.startswith(key + "-"):
            sub, weight = stem[len(key) + 1:].split("-")[:2]
            b64 = base64.b64encode(f.read_bytes()).decode()
            faces.append(f"@font-face{{font-family:'{fam}';font-style:normal;font-weight:{weight};font-display:swap;"
                         f"src:url(data:font/woff2;base64,{b64}) format('woff2');unicode-range:{RANGES[sub]}}}")
FONTS_CSS = "\n".join(faces)


def photo_url(p, w=1200):
    if not p:
        return None
    return p if p.startswith("http") else f"https://images.unsplash.com/photo-{p}?auto=format&fit=crop&w={w}&q=80"


# 1. Страница
tpl = (ROOT / "template.html").read_text(encoding="utf-8")
blob = json.dumps(data, ensure_ascii=False, indent=1).replace("</", "<\\/")
(OUT / "euphoria-card.html").write_text(tpl.replace("/*__FONTS__*/", FONTS_CSS).replace("/*__DATA__*/", blob), encoding="utf-8")

# 2. Прайс по шаблону Яндекс Бизнеса (О компании → Товары и услуги → Загрузить XLS)
cols = ["Категория", "Название", "Идентификатор", "Описание", "Короткое описание", "Цена",
        "Ссылка", "Фотография", "Популярный товар", "В наличии", "Количество", "Единицы измерения"]
wb = Workbook()
ws = wb.active
ws.title = "Прайс"
ws.append(cols)
for c in ws[1]:
    c.font = Font(bold=True)
yellow = PatternFill("solid", start_color="FFF2A8")
popular_left = 10
for s in data["services"]:
    pop = s.get("popular") and popular_left > 0
    popular_left -= bool(pop)
    short = s["short"] or (("Цена " + s["priceText"]) if "/" in s["priceText"] or s["priceFrom"] else "")
    ws.append([s["cat"], s["name"], s["id"], s["description"], short, s["price"],
               S.get("site") or None, photo_url(s.get("photo")), "да" if pop else "нет", "да", None, "услуга"])
    ws.cell(row=ws.max_row, column=8).fill = yellow  # фото пока примерные — заменить на фото салона
for i, w in enumerate([22, 46, 12, 64, 34, 10, 20, 34, 16, 11, 11, 18], start=1):
    ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True)
ws.freeze_panes = "A2"
wb.save(IMP / "price-yandex-business.xlsx")

# 3. Картинки: истории 1080×1920, логотип, обложка
story_tpl = (ROOT / "story_template.html").read_text(encoding="utf-8").replace("__FONTS__", FONTS_CSS)
brand_tpl = (ROOT / "brand_template.html").read_text(encoding="utf-8").replace("__FONTS__", FONTS_CSS)
try:
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 1080, "height": 1920})
        for st in data["stories"]:
            lines = "".join(f"<div><span>{html.escape(a)}</span><i></i><b>{html.escape(b)}</b></div>" for a, b in st.get("lines", []))
            bgf = st.get("bg")
            bg = ""
            if bgf:
                bp = ROOT / "photos" / ("stories/" + bgf if bgf.endswith(".jpg") else "services/" + bgf + ".jpg")
                bg = "data:image/jpeg;base64," + base64.b64encode(bp.read_bytes()).decode()
            page.set_content(story_tpl.replace("__BG__", bg).replace("__KICKER__", html.escape(st["kicker"])).replace("__TITLE__", html.escape(st["title"]))
                             .replace("__TEXT__", html.escape(st["text"])).replace("__LINES__", lines))
            page.evaluate("document.fonts.ready")
            page.wait_for_timeout(200)
            page.screenshot(path=str(IMP / "stories" / st["file"]))
        jobs = [
            ("logo-800x800.png", 800, 800,
             '<div class="frame" style="inset:40px;border-radius:50%"></div><div class="bt" style="font-size:52px">Beauty Salon</div>'
             '<div class="bn" style="font-size:118px;margin-top:6px">EUPHORIA</div>'),
            ("cover-1600x900.png", 1600, 900,
             '<img src="data:image/jpeg;base64,' + base64.b64encode((ROOT / "photos/salon/cover.jpg").read_bytes()).decode() + '" '
             'style="position:absolute;inset:0;width:100%;height:100%;object-fit:cover;object-position:60% 50%">'
             '<div style="position:absolute;inset:0;background:linear-gradient(90deg,rgba(11,11,12,.9) 0%,rgba(11,11,12,.78) 38%,rgba(11,11,12,.25) 68%,rgba(11,11,12,.05) 100%)"></div>'
             '<div class="frame" style="inset:36px;border-radius:24px"></div>'
             '<div style="position:absolute;left:110px;top:50%;transform:translateY(-50%);text-align:left;width:760px">'
             '<div class="bt" style="font-size:40px">Beauty Salon</div>'
             '<div class="bn" style="font-size:132px;margin-top:4px">EUPHORIA</div><div class="rule" style="width:220px;margin:36px 0 30px"></div>'
             f'<div class="tl" style="font-size:32px;color:#D6D2CD">{html.escape(S["tagline"])}</div>'
             '<div class="tl" style="font-size:28px;margin-top:20px;color:#D6D2CD">Реутов, Парковая ул., 8 · ежедневно 10–21</div></div>'),
        ]
        jobs.append(("../../server/bot-avatar-640x640.png", 640, 640,
                     '<div class="bt" style="font-size:40px">Beauty Salon</div><div class="bn" style="font-size:96px;margin-top:4px">EUPHORIA</div>'
                     '<div class="tl" style="font-size:30px;margin-top:26px">заявки</div>'))
        (OUT / "server").mkdir(exist_ok=True)
        for name, w, h, body in jobs:
            page.set_viewport_size({"width": w, "height": h})
            page.set_content(brand_tpl.replace("__W__", str(w)).replace("__H__", str(h)).replace("__BODY__", body))
            page.evaluate("document.fonts.ready")
            page.wait_for_timeout(200)
            page.screenshot(path=str(IMP / "brand-images" / name))
        browser.close()
except Exception as e:  # картинки не обязательны для остальной сборки
    print("Картинки не собраны:", e)

# 4. Тексты для кабинета
site_note = S["site"] or "ссылка на сайт или соцсети салона"
lines = ["ТЕКСТЫ ДЛЯ КАБИНЕТА ЯНДЕКС БИЗНЕСА", "=" * 40, "",
         "ОПИСАНИЕ ОРГАНИЗАЦИИ (О компании → Описание)", S["description"], "",
         "ИСТОРИИ (Истории → Добавить историю). Картинки в папке «stories».", ""]
for i, st in enumerate(data["stories"], start=1):
    lines += [f"История {i}: файл {st['file']}", f"  Название истории: {st['caption']}",
              f"  Текст: {st['title']}. {st['text']}", f"  Кнопка: {st['cta']} → {site_note}", ""]
lines += ["ПРОМО-СТРОКА (только с подпиской на продвижение)", data["offer"]["promoLine"], "",
          "ОСОБЕННОСТИ (О компании → Особенности)"] + [f"  {k}: {v}" for k, v in data["features"]]
(IMP / "texts-for-cabinet.txt").write_text("\n".join(lines), encoding="utf-8")

# 5. Сервер заявок (локальный, без облаков), инструкция, исходники
SRV = OUT / "server"
(SRV / "public").mkdir(parents=True, exist_ok=True)
for f in ["server.py", "start.bat", "start.command"]:
    shutil.copy(ROOT / "server" / f, SRV / f)
shutil.copy(OUT / "euphoria-card.html", SRV / "public" / "index.html")

# 5а. Версия для обновления серверов: server/manifest.json в репозитории (server.py сверяет его с файлом VERSION)
import hashlib
import sys
from datetime import datetime, timedelta, timezone
(ROOT / "server" / "public").mkdir(exist_ok=True)
shutil.copy(OUT / "euphoria-card.html", ROOT / "server" / "public" / "index.html")
files = {f: hashlib.sha256((ROOT / "server" / f).read_bytes()).hexdigest() for f in ["server.py", "public/index.html"]}
version = hashlib.sha256(json.dumps(files, sort_keys=True).encode()).hexdigest()[:12]
man_path = ROOT / "server" / "manifest.json"
old = json.loads(man_path.read_text(encoding="utf-8")) if man_path.exists() else {}
note = " ".join(sys.argv[1:]).strip()
if old.get("version") == version:
    man = dict(old, note=note or old.get("note", ""))
else:
    man = {"version": version, "built": datetime.now(timezone(timedelta(hours=3))).strftime("%d.%m.%Y %H:%M"), "note": note, "files": files}
man_path.write_text(json.dumps(man, ensure_ascii=False, indent=2), encoding="utf-8")
(SRV / "VERSION").write_text(version, encoding="utf-8")
shutil.copy(ROOT / "ИНСТРУКЦИЯ.txt", OUT / "INSTRUKCIYA.txt")
src = OUT / "sources"
src.mkdir()
for f in ["data.json", "gen_data.py", "build.py", "template.html", "story_template.html", "brand_template.html"]:
    shutil.copy(ROOT / f, src / f)
shutil.copytree(ROOT / "fonts", src / "fonts")
print("Готово:", OUT, "·", len(data["services"]), "услуг")

# 6. Архив (латинские имена — встроенный архиватор Windows не портит их)
import zipfile
zp = OUT.parent / "euphoria.zip"
zp.unlink(missing_ok=True)
with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
    for f in sorted(OUT.rglob("*")):
        z.write(f, "euphoria/" + f.relative_to(OUT).as_posix())
print("Архив:", zp)
